"""Main application routes: dashboard, tenants, payments, reports."""
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from io import BytesIO

from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    jsonify,
    send_file,
    abort,
)
from flask_login import login_required
from sqlalchemy import func
from sqlalchemy.exc import SQLAlchemyError

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from models import db, Tenant, Payment

main_bp = Blueprint("main", __name__)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _parse_decimal(value, default=Decimal("0")) -> Decimal:
    if value is None or value == "":
        return default
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError(f"Invalid number: {value!r}")


def _parse_date(value, default=None):
    if not value:
        return default
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        raise ValueError(f"Invalid date: {value!r}")


def _parse_int(value, default=0, min_v=None, max_v=None) -> int:
    try:
        n = int(value)
    except (TypeError, ValueError):
        return default
    if min_v is not None:
        n = max(min_v, n)
    if max_v is not None:
        n = min(max_v, n)
    return n


def _get_overdue_tenants():
    """Return list of overdue Tenant objects (active, balance > 0, past due day)."""
    today = date.today()
    tenants = Tenant.query.filter(
        Tenant.status == "active",
        Tenant.pending_balance > 0,
        Tenant.due_day < today.day,
    ).all()
    return tenants


def _get_unpaid_tenants():
    return Tenant.query.filter(
        Tenant.status == "active", Tenant.pending_balance > 0
    ).all()


def _get_next_due_date(due_day: int, current_date: date = None) -> date:
    """Calculate the next due date based on due day of month."""
    if current_date is None:
        current_date = date.today()
    
    if current_date.day <= due_day:
        # Due date is in the same month
        return date(current_date.year, current_date.month, due_day)
    else:
        # Due date is in the next month
        if current_date.month == 12:
            return date(current_date.year + 1, 1, due_day)
        else:
            return date(current_date.year, current_date.month + 1, due_day)


def _get_upcoming_urgent_tenants():
    """Return list of tenants with upcoming (7 days) or urgent (3 days) dues."""
    today = date.today()
    tenants = []
    
    for tenant in Tenant.query.filter(
        Tenant.status == "active",
        Tenant.pending_balance > 0
    ).all():
        next_due_date = _get_next_due_date(tenant.due_day, today)
        days_remaining = (next_due_date - today).days
        
        if days_remaining <= 7 and days_remaining >= 0:
            urgency = "urgent" if days_remaining <= 3 else "upcoming"
            tenants.append({
                "id": tenant.id,
                "nickname": tenant.nickname,
                "full_name": tenant.full_name,
                "room_number": tenant.room_number,
                "due_date": next_due_date.strftime("%b %d, %Y"),
                "days_remaining": days_remaining,
                "urgency": urgency,
                "pending_balance": float(tenant.pending_balance or 0)
            })
    
    # Sort: urgent first, then by days remaining
    tenants.sort(key=lambda x: (x["urgency"] != "urgent", x["days_remaining"]))
    return tenants


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

@main_bp.route("/")
@login_required
def dashboard():
    try:
        total_tenants = Tenant.query.count()
        active_tenants = Tenant.query.filter_by(status="active").count()

        # Monthly income = payments in current month
        today = date.today()
        monthly_income = (
            db.session.query(func.coalesce(func.sum(Payment.amount), 0))
            .filter(
                func.extract("year", Payment.payment_date) == today.year,
                func.extract("month", Payment.payment_date) == today.month,
            )
            .scalar()
            or 0
        )

        unpaid = _get_unpaid_tenants()
        overdue = _get_overdue_tenants()

        recent_payments = (
            Payment.query.order_by(Payment.payment_date.desc(), Payment.id.desc())
            .limit(5)
            .all()
        )

        response = render_template(
            "dashboard.html",
            total_tenants=total_tenants,
            active_tenants=active_tenants,
            monthly_income=float(monthly_income),
            unpaid_count=len(unpaid),
            overdue_count=len(overdue),
            overdue_tenants=overdue,
            recent_payments=recent_payments,
        )
        response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
        return response
    except Exception as e:
        app.logger.error(f"Error loading dashboard: {e}")
        flash("Error loading dashboard data. Please try again.", "error")
        return render_template("dashboard.html", total_tenants=0, active_tenants=0, monthly_income=0, unpaid_count=0, overdue_count=0, overdue_tenants=[], recent_payments=[])


# ---------------------------------------------------------------------------
# Notifications API (used for live refresh of overdue badge)
# ---------------------------------------------------------------------------

@main_bp.route("/api/notifications")
@login_required
def api_notifications():
    overdue = _get_overdue_tenants()
    unpaid = _get_unpaid_tenants()
    upcoming_urgent = _get_upcoming_urgent_tenants()
    return jsonify(
        {
            "overdue_count": len(overdue),
            "unpaid_count": len(unpaid),
            "notification_count": len(upcoming_urgent),
            "overdue_tenants": [
                {
                    "id": t.id,
                    "nickname": t.nickname,
                    "full_name": t.full_name,
                    "room_number": t.room_number,
                    "pending_balance": float(t.pending_balance or 0),
                }
                for t in overdue
            ],
            "upcoming_urgent_tenants": upcoming_urgent,
        }
    )


# ---------------------------------------------------------------------------
# Tenants CRUD
# ---------------------------------------------------------------------------

@main_bp.route("/tenants")
@login_required
def tenants_list():
    try:
        search = (request.args.get("q") or "").strip()
        status_filter = request.args.get("status", "")

        query = Tenant.query
        if search:
            like = f"%{search}%"
            query = query.filter(
                db.or_(
                    Tenant.nickname.ilike(like),
                    Tenant.full_name.ilike(like),
                    Tenant.room_number.ilike(like),
                )
            )
        if status_filter in ("active", "inactive"):
            query = query.filter(Tenant.status == status_filter)

        tenants = query.order_by(Tenant.room_number.asc(), Tenant.nickname.asc()).all()
        response = render_template(
            "tenants.html", tenants=tenants, search=search, status_filter=status_filter
        )
        response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
        return response
    except Exception as e:
        app.logger.error(f"Error loading tenants: {e}")
        flash("Error loading tenant data. Please try again.", "error")
        return render_template("tenants.html", tenants=[], search="", status_filter="")


@main_bp.route("/tenants/new", methods=["GET", "POST"])
@login_required
def tenant_create():
    if request.method == "POST":
        try:
            tenant = Tenant(
                nickname=(request.form.get("nickname") or "").strip(),
                full_name=(request.form.get("full_name") or "").strip(),
                room_number=(request.form.get("room_number") or "").strip(),
                contact_number=(request.form.get("contact_number") or "").strip() or None,
                move_in_date=_parse_date(request.form.get("move_in_date"), date.today()),
                monthly_rent=_parse_decimal(request.form.get("monthly_rent")),
                due_day=_parse_int(request.form.get("due_day"), 1, 1, 31),
                pending_balance=_parse_decimal(request.form.get("pending_balance")),
                status=request.form.get("status", "active"),
            )

            if not tenant.nickname or not tenant.full_name or not tenant.room_number:
                flash("Nickname, full name, and room number are required.", "error")
                return render_template("tenant_form.html", tenant=None, form=request.form)

            if Tenant.query.filter_by(nickname=tenant.nickname).first():
                flash(f"Nickname '{tenant.nickname}' is already taken.", "error")
                return render_template("tenant_form.html", tenant=None, form=request.form)

            db.session.add(tenant)
            db.session.commit()
            flash(f"Tenant '{tenant.nickname}' added successfully.", "success")
            return redirect(url_for("main.tenants_list"))
        except ValueError as e:
            flash(str(e), "error")
            return render_template("tenant_form.html", tenant=None, form=request.form)
        except SQLAlchemyError as e:
            db.session.rollback()
            flash(f"Database error: {e}", "error")
            return render_template("tenant_form.html", tenant=None, form=request.form)

    return render_template("tenant_form.html", tenant=None, form={})


@main_bp.route("/tenants/<int:tenant_id>")
@login_required
def tenant_detail(tenant_id: int):
    try:
        tenant = Tenant.query.get_or_404(tenant_id)
        payments = tenant.payments.all()
        total_paid = sum((float(p.amount) for p in payments), 0.0)
        response = render_template(
            "tenant_detail.html",
            tenant=tenant,
            payments=payments,
            total_paid=total_paid,
        )
        response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
        return response
    except Exception as e:
        app.logger.error(f"Error loading tenant detail: {e}")
        flash("Error loading tenant data. Please try again.", "error")
        return redirect(url_for("main.tenants_list"))


@main_bp.route("/tenants/<int:tenant_id>/edit", methods=["GET", "POST"])
@login_required
def tenant_edit(tenant_id: int):
    tenant = Tenant.query.get_or_404(tenant_id)

    if request.method == "POST":
        try:
            new_nickname = (request.form.get("nickname") or "").strip()
            if new_nickname != tenant.nickname:
                if Tenant.query.filter_by(nickname=new_nickname).first():
                    flash(f"Nickname '{new_nickname}' is already taken.", "error")
                    return render_template(
                        "tenant_form.html", tenant=tenant, form=request.form
                    )

            tenant.nickname = new_nickname
            tenant.full_name = (request.form.get("full_name") or "").strip()
            tenant.room_number = (request.form.get("room_number") or "").strip()
            tenant.contact_number = (request.form.get("contact_number") or "").strip() or None
            tenant.move_in_date = _parse_date(
                request.form.get("move_in_date"), tenant.move_in_date
            )
            tenant.monthly_rent = _parse_decimal(request.form.get("monthly_rent"))
            tenant.due_day = _parse_int(request.form.get("due_day"), 1, 1, 31)
            tenant.pending_balance = _parse_decimal(request.form.get("pending_balance"))
            tenant.status = request.form.get("status", "active")

            if not tenant.nickname or not tenant.full_name or not tenant.room_number:
                flash("Nickname, full name, and room number are required.", "error")
                return render_template("tenant_form.html", tenant=tenant, form=request.form)

            db.session.commit()
            flash(f"Tenant '{tenant.nickname}' updated.", "success")
            return redirect(url_for("main.tenant_detail", tenant_id=tenant.id))
        except ValueError as e:
            flash(str(e), "error")
        except SQLAlchemyError as e:
            db.session.rollback()
            flash(f"Database error: {e}", "error")

    return render_template("tenant_form.html", tenant=tenant, form={})


@main_bp.route("/tenants/<int:tenant_id>/delete", methods=["POST"])
@login_required
def tenant_delete(tenant_id: int):
    tenant = Tenant.query.get_or_404(tenant_id)
    try:
        db.session.delete(tenant)
        db.session.commit()
        flash(f"Tenant '{tenant.nickname}' deleted.", "success")
    except SQLAlchemyError as e:
        db.session.rollback()
        flash(f"Could not delete tenant: {e}", "error")
    return redirect(url_for("main.tenants_list"))


# ---------------------------------------------------------------------------
# Payments
# ---------------------------------------------------------------------------

@main_bp.route("/payments")
@login_required
def payments_list():
    try:
        tenant_id = request.args.get("tenant_id", type=int)
        query = Payment.query
        if tenant_id:
            query = query.filter_by(tenant_id=tenant_id)
        payments = query.order_by(
            Payment.payment_date.desc(), Payment.id.desc()
        ).all()
        tenants = Tenant.query.order_by(Tenant.nickname.asc()).all()
        total = sum((float(p.amount) for p in payments), 0.0)
        response = render_template(
            "payments.html",
            payments=payments,
            tenants=tenants,
            tenant_id=tenant_id,
            total=total,
        )
        response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
        return response
    except Exception as e:
        app.logger.error(f"Error loading payments: {e}")
        flash("Error loading payment data. Please try again.", "error")
        return render_template("payments.html", payments=[], tenants=[], tenant_id=None, total=0)


@main_bp.route("/payments/new", methods=["GET", "POST"])
@login_required
def payment_create():
    try:
        tenants = Tenant.query.order_by(Tenant.nickname.asc()).all()
    except Exception as e:
        app.logger.error(f"Error fetching tenants: {e}")
        flash("Error loading tenant data. Please try again.", "error")
        return redirect(url_for("main.payments_list"))
    
    preselect_id = request.args.get("tenant_id", type=int)

    if request.method == "POST":
        try:
            # Validate tenant selection
            tenant_id = _parse_int(request.form.get("tenant_id"), 0)
            if not tenant_id:
                flash("Please select a tenant.", "error")
                return render_template(
                    "payment_form.html",
                    tenants=tenants,
                    form=request.form,
                    preselect_id=preselect_id,
                )
            
            tenant = Tenant.query.get(tenant_id)
            if not tenant:
                flash("Please choose a valid tenant.", "error")
                return render_template(
                    "payment_form.html",
                    tenants=tenants,
                    form=request.form,
                    preselect_id=preselect_id,
                )

            # Validate amount
            amount = _parse_decimal(request.form.get("amount"))
            if amount is None or amount <= 0:
                flash("Payment amount must be greater than zero.", "error")
                return render_template(
                    "payment_form.html",
                    tenants=tenants,
                    form=request.form,
                    preselect_id=preselect_id,
                )

            # Validate payment date
            payment_date = _parse_date(request.form.get("payment_date"), date.today())
            if not payment_date:
                payment_date = date.today()

            # Create payment
            payment = Payment(
                tenant_id=tenant.id,
                amount=amount,
                payment_date=payment_date,
                notes=(request.form.get("notes") or "").strip() or None,
            )

            # Reduce pending balance (allow going into credit = negative)
            old_balance = tenant.pending_balance or 0
            tenant.pending_balance = old_balance - amount

            db.session.add(payment)
            db.session.commit()
            
            app.logger.info(f"Payment recorded: tenant={tenant.nickname}, amount={amount}, old_balance={old_balance}, new_balance={tenant.pending_balance}")
            
            flash(
                f"Payment of {amount} recorded for {tenant.nickname}.",
                "success",
            )
            return redirect(url_for("main.payments_list"))
        except ValueError as e:
            app.logger.error(f"Validation error: {e}")
            flash(str(e), "error")
        except SQLAlchemyError as e:
            db.session.rollback()
            app.logger.error(f"Database error: {e}")
            flash(f"Database error: {e}", "error")
        except Exception as e:
            db.session.rollback()
            app.logger.error(f"Unexpected error: {e}")
            flash(f"An error occurred: {e}", "error")

    return render_template(
        "payment_form.html",
        tenants=tenants,
        form={},
        preselect_id=preselect_id,
    )


@main_bp.route("/payments/<int:payment_id>/delete", methods=["POST"])
@login_required
def payment_delete(payment_id: int):
    payment = Payment.query.get_or_404(payment_id)
    try:
        tenant = payment.tenant
        if tenant:
            tenant.pending_balance = (tenant.pending_balance or 0) + payment.amount
        db.session.delete(payment)
        db.session.commit()
        flash("Payment removed and balance restored.", "success")
    except SQLAlchemyError as e:
        db.session.rollback()
        flash(f"Could not delete payment: {e}", "error")
    return redirect(url_for("main.payments_list"))


# ---------------------------------------------------------------------------
# Reports & Excel export
# ---------------------------------------------------------------------------


@main_bp.route("/reports/analytics")
@login_required
def reports_analytics():
    """Analytics API endpoint for monthly/yearly data aggregation."""
    try:
        # Get filter parameters
        filter_type = request.args.get("filter_type", "monthly")  # monthly or yearly
        year = request.args.get("year", date.today().year)
        month = request.args.get("month", date.today().month)
        
        # Validate inputs
        try:
            year = int(year)
            month = int(month) if month else None
        except (ValueError, TypeError):
            year = date.today().year
            month = date.today().month
        
        # Base query for payments
        payments_query = db.session.query(
            func.extract("year", Payment.payment_date).label("year"),
            func.extract("month", Payment.payment_date).label("month"),
            func.sum(Payment.amount).label("total_amount"),
            func.count(Payment.id).label("payment_count")
        ).group_by(
            func.extract("year", Payment.payment_date),
            func.extract("month", Payment.payment_date)
        ).order_by(
            func.extract("year", Payment.payment_date),
            func.extract("month", Payment.payment_date)
        )
        
        # Apply filters
        if filter_type == "monthly":
            payments_query = payments_query.filter(
                func.extract("year", Payment.payment_date) == year,
                func.extract("month", Payment.payment_date) == month
            )
        elif filter_type == "yearly":
            payments_query = payments_query.filter(
                func.extract("year", Payment.payment_date) == year
            )
        
        payments_data = payments_query.all()
        
        # Process data for response
        analytics_data = []
        for row in payments_data:
            analytics_data.append({
                "year": int(row.year) if row.year else None,
                "month": int(row.month) if row.month else None,
                "total_amount": float(row.total_amount or 0),
                "payment_count": row.payment_count or 0
            })
        
        # Calculate total income for the selected period
        total_income_period = sum(p["total_amount"] for p in analytics_data)
        total_payments_period = sum(p["payment_count"] for p in analytics_data)
        
        # Get tenant statistics
        tenant_stats = {
            "total_tenants": Tenant.query.count(),
            "active_tenants": Tenant.query.filter_by(status="active").count(),
            "inactive_tenants": Tenant.query.filter_by(status="inactive").count()
        }
        
        # Get balance statistics (current outstanding balance - not filtered by period)
        balance_stats = {
            "total_balance": float(
                db.session.query(func.coalesce(func.sum(Tenant.pending_balance), 0))
                .filter(Tenant.status == "active", Tenant.pending_balance > 0)
                .scalar() or 0
            ),
            "overdue_count": len(_get_overdue_tenants()),
            "unpaid_count": len(_get_unpaid_tenants())
        }
        
        return jsonify({
            "success": True,
            "filter_type": filter_type,
            "year": year,
            "month": month,
            "payments": analytics_data,
            "total_income_period": total_income_period,
            "total_payments_period": total_payments_period,
            "tenant_stats": tenant_stats,
            "balance_stats": balance_stats
        })
    except Exception as e:
        return jsonify({
            "success": False,
            "error": str(e)
        }), 500

@main_bp.route("/reports")
@login_required
def reports():
    try:
        total_tenants = Tenant.query.count()
        active_tenants = Tenant.query.filter_by(status="active").count()
        inactive_tenants = total_tenants - active_tenants

        total_income = (
            db.session.query(func.coalesce(func.sum(Payment.amount), 0)).scalar() or 0
        )
        today = date.today()
        monthly_income = (
            db.session.query(func.coalesce(func.sum(Payment.amount), 0))
            .filter(
                func.extract("year", Payment.payment_date) == today.year,
                func.extract("month", Payment.payment_date) == today.month,
            )
            .scalar()
            or 0
        )

        total_balance = (
            db.session.query(func.coalesce(func.sum(Tenant.pending_balance), 0))
            .filter(Tenant.status == "active", Tenant.pending_balance > 0)
            .scalar()
            or 0
        )

        overdue_tenants = _get_overdue_tenants()
        unpaid_tenants = _get_unpaid_tenants()

        response = render_template(
            "reports.html",
            total_tenants=total_tenants,
            active_tenants=active_tenants,
            inactive_tenants=inactive_tenants,
            total_income=float(total_income),
            monthly_income=float(monthly_income),
            total_balance=float(total_balance),
            overdue_tenants=overdue_tenants,
            unpaid_tenants=unpaid_tenants,
        )
        response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
        return response
    except Exception as e:
        app.logger.error(f"Error loading reports: {e}")
        flash("Error loading reports data. Please try again.", "error")
        return render_template(
            "reports.html",
            total_tenants=0,
            active_tenants=0,
            inactive_tenants=0,
            total_income=0,
            monthly_income=0,
            total_balance=0,
            overdue_tenants=[],
            unpaid_tenants=[],
        )


def _style_header(ws, headers):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="5C4033")  # deep brown
    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(12, max_len + 2), 40)


@main_bp.route("/export/tenants.xlsx")
@login_required
def export_tenants():
    tenants = Tenant.query.order_by(Tenant.room_number.asc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Tenants"

    headers = [
        "ID", "Nickname", "Full Name", "Room", "Contact",
        "Move-in Date", "Monthly Rent", "Due Day",
        "Pending Balance", "Status", "Created At",
    ]
    _style_header(ws, headers)

    for t in tenants:
        ws.append([
            t.id,
            t.nickname,
            t.full_name,
            t.room_number,
            t.contact_number or "",
            t.move_in_date.isoformat() if t.move_in_date else "",
            float(t.monthly_rent or 0),
            t.due_day,
            float(t.pending_balance or 0),
            t.status,
            t.created_at.strftime("%Y-%m-%d %H:%M") if t.created_at else "",
        ])

    _auto_width(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"tenants_{date.today().isoformat()}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@main_bp.route("/export/payments.xlsx")
@login_required
def export_payments():
    payments = (
        Payment.query.join(Tenant)
        .order_by(Payment.payment_date.desc(), Payment.id.desc())
        .all()
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Payments"

    headers = ["ID", "Tenant Nickname", "Tenant Full Name", "Room", "Amount", "Payment Date", "Notes"]
    _style_header(ws, headers)

    for p in payments:
        ws.append([
            p.id,
            p.tenant.nickname if p.tenant else "",
            p.tenant.full_name if p.tenant else "",
            p.tenant.room_number if p.tenant else "",
            float(p.amount or 0),
            p.payment_date.isoformat() if p.payment_date else "",
            p.notes or "",
        ])

    _auto_width(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"payments_{date.today().isoformat()}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
